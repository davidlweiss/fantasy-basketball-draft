"""Read the projections workbook and write data/players.json.

players.json = {"p": [[name, rank, last_year_rank, adp, do_not_draft, fgm, fga, ftm, fta,
                       3pm, pts, reb, ast, stl, blk, to], ...],
                "t": {name: "3-4", ...}}   # starting round targets from the 'real draft' tab
Run:  python src/extract_data.py
"""
import json, os, re, warnings
import openpyxl
warnings.filterwarnings("ignore")
ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
CFG = json.load(open(os.path.join(ROOT, "config.json")))
wv = openpyxl.load_workbook(os.path.join(ROOT, CFG["xlsx"]), data_only=True)

def made_att(text):                      # '0.579(10.9/18.9)' -> (10.9, 18.9)
    m = re.match(r"([\d.]+)\(([\d.]+)/([\d.]+)\)", str(text))
    return (float(m[2]), float(m[3])) if m else (None, None)

num = lambda v: v if isinstance(v, (int, float)) else None
players = []
for r in wv["2026-2027 projections"].iter_rows(min_row=2, values_only=True):
    if not r[1] or r[2] is None:
        continue
    fgm, fga = made_att(r[7]); ftm, fta = made_att(r[8])
    if fgm is None:
        continue
    players.append([r[1], r[2], num(r[3]), num(r[4]), int(bool(r[6])), fgm, fga, ftm, fta,
                    r[9], r[10], r[11], r[12], r[13], r[14], r[15]])

# starting target lists: the 'real draft' tab, column L, grouped by 'Rds1-2' style headers
cur, groups = None, {}
for (v,) in wv["real draft"].iter_rows(min_row=2, max_row=109, min_col=12, max_col=12, values_only=True):
    if not v:
        continue
    m = re.match(r"Rds?\s*(\d+)-(\d+)", v, re.I)
    if m:
        cur = f"{m[1]}-{m[2]}"; continue
    groups.setdefault(cur or "1-2", []).append(v)      # first, unlabeled block = rounds 1-2
norm = lambda s: re.sub(r"[^a-z0-9]", "", s.lower())
by = {norm(p[0]): p[0] for p in players}
targets, unmatched = {}, []
for rd, names in groups.items():
    for n in names:
        if norm(n) in by:
            targets.setdefault(by[norm(n)], rd)           # earliest round wins
        else:
            unmatched.append((rd, n))
os.makedirs(os.path.join(ROOT, "data"), exist_ok=True)
json.dump({"p": players, "t": targets}, open(os.path.join(ROOT, "data", "players.json"), "w"), separators=(",", ":"))
print(f"{len(players)} players, {len(targets)} starting targets matched")
print("not in the 2026-27 pool (ignored):", unmatched)
